import asyncio
import functools
import io
from datetime import date, datetime, time

from aiogram import Bot
from aiogram.types import BufferedInputFile, CallbackQuery, Message
from aiogram_dialog import Dialog, DialogManager, StartMode, Window
from aiogram_dialog.widgets.input import MessageInput
from aiogram_dialog.widgets.kbd import Button, Calendar, Column, Row, ScrollingGroup, Select
from aiogram_dialog.widgets.text import Const, Format
from dishka import FromDishka
from dishka.integrations.aiogram_dialog import inject
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from quizzi.application.bot.shared_dialogs.states import SharedCreateTestSG, SharedTestsSG
from quizzi.infrastructure.database.dao.group import GroupDAO
from quizzi.infrastructure.database.dao.test import TestDAO
from quizzi.infrastructure.database.repo.test import TestRepository
from quizzi.infrastructure.database.repo.test_attempt import TestAttemptRepository
from quizzi.infrastructure.utils.config import Config
from quizzi.infrastructure.utils.qr_generator import generate_qr_bytes
from quizzi.infrastructure.utils.test_id_to_hash import encode_id
from quizzi.infrastructure.utils.timezone import to_msk


@inject
async def get_tests_data(test_dao: FromDishka[TestDAO], **_kwargs):
    tests = await test_dao.get_all()
    
    return {
        "tests": [
            (f"{'🟢' if t.is_active else '🔴'} {t.title}", t.id)
            for t in tests
        ],
        "count": len(tests),
    }


async def on_test_selected(_callback: CallbackQuery, _widget: Select, manager: DialogManager, item_id: str):
    manager.dialog_data["selected_test_id"] = int(item_id)
    await manager.switch_to(SharedTestsSG.test_detail)


@inject
async def get_test_detail(test_dao: FromDishka[TestDAO], test_repo: FromDishka[TestRepository], dialog_manager: DialogManager, **_kwargs):
    test_id = dialog_manager.dialog_data.get("selected_test_id")
    
    if not test_id:
        return {
            "test_info": "Тест не найден",
            "is_active": False,
            "button_text": "◀️ Назад",
            "results_button_text": "👁 Показать результаты",
        }
    
    test = await test_dao.get_by_id(test_id)
    questions_count = await test_repo.count_questions_in_test(test_id)
    
    if not test:
        return {
            "test_info": "Тест не найден",
            "is_active": False,
            "button_text": "◀️ Назад",
            "results_button_text": "👁 Показать результаты",
        }
    
    status = "🟢 Активен" if test.is_active else "🔴 Деактивирован"
    password_str = f"🔒 {test.password}" if test.password else "🔓 Без пароля"
    attempts_str = f"🔄 {test.attempts}" if test.attempts else "♾️ Без ограничений"
    time_limit_str = f"⏱️ {test.time_limit // 60} мин" if test.time_limit else "⏱️ Без лимита"
    expires_str = f"📅 {to_msk(test.expires_at).strftime('%d.%m.%Y %H:%M')}" if test.expires_at else "📅 Без срока"
    group_str = f"🎓 Группа {test.for_group}" if test.for_group else "👥 Для всех"
    results_str = "👁 Результаты видны" if test.are_results_viewable else "🔒 Результаты скрыты"
    
    test_info = (
        f"<b>📝 Информация о тесте</b>\n\n"
        f"<b>Название:</b>\n<blockquote>{test.title}</blockquote>\n"
        f"<b>Описание:</b>\n<blockquote>{test.description or '—'}</blockquote>\n\n"
        f"<b>Статус:</b> {status}\n"
        f"<b>Вопросов:</b> {questions_count}\n"
        f"<b>Пароль:</b> {password_str}\n"
        f"<b>Попытки:</b> {attempts_str}\n"
        f"<b>Время:</b> {time_limit_str}\n"
        f"<b>Срок:</b> {expires_str}\n"
        f"<b>Группа:</b> {group_str}\n"
        f"<b>Видимость:</b> {results_str}\n\n"
        f"<b>Создан:</b> {to_msk(test.created_at).strftime('%d.%m.%Y %H:%M') if test.created_at else '—'}"
    )
    
    button_text = "🔴 Деактивировать" if test.is_active else "🟢 Активировать"
    results_button_text = "🔒 Скрыть результаты" if test.are_results_viewable else "👁 Показать результаты"
    
    return {
        "test_info": test_info,
        "is_active": test.is_active,
        "button_text": button_text,
        "results_button_text": results_button_text,
    }


@inject
async def on_toggle_active(_callback: CallbackQuery, _button: Button, manager: DialogManager, test_dao: FromDishka[TestDAO]):
    test_id = manager.dialog_data.get("selected_test_id")
    if not test_id:
        await _callback.answer("❌ Тест не найден")
        return
    
    test = await test_dao.get_by_id(test_id)
    
    if test:
        await test_dao.update(test_id, is_active=not test.is_active)
        action = "деактивирован" if test.is_active else "активирован"
        await _callback.answer(f"✅ Тест {action}")
        await manager.switch_to(SharedTestsSG.test_detail)


@inject
async def on_toggle_results_viewable(_callback: CallbackQuery, _button: Button, manager: DialogManager, test_dao: FromDishka[TestDAO]):
    test_id = manager.dialog_data.get("selected_test_id")
    if not test_id:
        await _callback.answer("❌ Тест не найден")
        return
    
    test = await test_dao.get_by_id(test_id)
    
    if test:
        await test_dao.update(test_id, are_results_viewable=not test.are_results_viewable)
        action = "скрыты" if test.are_results_viewable else "видны"
        await _callback.answer(f"✅ Результаты теперь {action}")
        await manager.switch_to(SharedTestsSG.test_detail)


async def on_back_to_list(_callback: CallbackQuery, _button: Button, manager: DialogManager):
    await manager.switch_to(SharedTestsSG.tests_list)


async def on_statistics(_callback: CallbackQuery, _button: Button, manager: DialogManager):
    await manager.switch_to(SharedTestsSG.statistics)


@inject
async def get_statistics_data(
    dialog_manager: DialogManager,
    attempt_repo: FromDishka[TestAttemptRepository],
    **_kwargs
):
    test_id = dialog_manager.dialog_data.get("selected_test_id")
    
    if not test_id:
        return {"attempts": [], "count": 0}
    
    attempts_with_users = await attempt_repo.get_test_attempts_with_users(test_id)
    
    results = []
    for attempt, user_name in attempts_with_users:
        status = "✅" if attempt.is_passed else "❌"
        finished_at_msk = to_msk(attempt.finished_at)
        date_str = finished_at_msk.strftime("%d.%m.%Y %H:%M") if finished_at_msk else ""
        results.append((f"{status} {user_name} — {attempt.score}% ({date_str})", attempt.id))
    
    return {
        "attempts": results,
        "count": len(results),
    }


async def on_attempt_selected(_callback: CallbackQuery, _widget: Select, manager: DialogManager, item_id: str):
    manager.dialog_data["selected_attempt_id"] = int(item_id)
    await manager.switch_to(SharedTestsSG.attempt_detail)


async def on_back_to_statistics(_callback: CallbackQuery, _button: Button, manager: DialogManager):
    await manager.switch_to(SharedTestsSG.statistics)


@inject
async def get_attempt_detail(
    dialog_manager: DialogManager,
    attempt_repo: FromDishka[TestAttemptRepository],
    test_repo: FromDishka[TestRepository],
    **_kwargs
):
    attempt_id = dialog_manager.dialog_data.get("selected_attempt_id")
    
    if not attempt_id:
        return {"attempt_info": "❌ Результат не найден"}
    
    attempt, answers = await attempt_repo.get_attempt_with_answers(attempt_id)
    
    if not attempt:
        return {"attempt_info": "❌ Результат не найден"}
    
    status = "✅ Пройден" if attempt.is_passed else "❌ Не пройден"
    finished_at_msk = to_msk(attempt.finished_at)
    date_str = finished_at_msk.strftime("%d.%m.%Y %H:%M") if finished_at_msk else "—"
    
    lines = [
        "<b>📊 Результат прохождения</b>\n",
        f"📈 <b>Результат:</b> {attempt.score}%",
        f"📅 <b>Дата:</b> {date_str}",
        f"🏆 <b>Статус:</b> {status}\n",
        "<b>📋 Ответы:</b>\n",
    ]
    
    question_ids = [answer.question_id for answer in answers]
    questions_map = await test_repo.get_questions_with_options_by_ids(question_ids)
    
    for i, answer in enumerate(answers, 1):
        question_data = questions_map.get(answer.question_id)
        if not question_data:
            continue
        
        question, options = question_data
        correct_options = [opt for opt in options if opt.is_correct]
        correct_texts = [opt.text for opt in correct_options]
        
        status_icon = "✅" if answer.is_correct else "❌"
        
        user_answer = answer.text_answer or ""
        if "|" in user_answer:
            user_answer = ", ".join(user_answer.split("|"))
        
        lines.append(f"{status_icon} <b>Вопрос {i}</b>")
        lines.append(f"<blockquote>{question.text}</blockquote>")
        lines.append(f"👤 <i>Ответ:</i> {user_answer or '—'}")
        lines.append(f"✓ <i>Правильно:</i> {', '.join(correct_texts)}\n")
    
    return {"attempt_info": "\n".join(lines)}


async def on_export_stats(_callback: CallbackQuery, _button: Button, manager: DialogManager) -> None:
    await manager.switch_to(SharedTestsSG.export_select_group)


@inject
async def on_share_test(_callback: CallbackQuery, _button: Button, manager: DialogManager, config: FromDishka[Config], bot_inst: FromDishka[Bot]):
    test_id = manager.dialog_data.get("selected_test_id")
    
    if not test_id:
        return {
            "share_link": "Ошибка: тест не найден"
        }
    
    test_hash = encode_id(
        test_id, 
        config.security.encode_key,
        config.security.encoded_string_length
    )
    
    bot_info = await bot_inst.get_me()
    bot_username = bot_info.username or "your_bot"
    share_link = f"https://t.me/{bot_username}?start={test_hash}"
    
    loop = asyncio.get_running_loop()
    qr_bytes = await loop.run_in_executor(
        None,
        functools.partial(generate_qr_bytes, share_link)
    )

    assert _callback.message is not None

    await _callback.message.answer_photo(
        photo=BufferedInputFile(qr_bytes, filename="qr.png"),
        caption=f"<b>🔗 Поделиться тестом</b>\n\n📎 <b>Ссылка на тест:</b>\n<code>{share_link}</code>\n\n💡 Отправьте эту ссылку или QR-код пользователям для прохождения теста"
    )


async def on_edit_menu(_callback: CallbackQuery, _button: Button, manager: DialogManager):
    await manager.switch_to(SharedTestsSG.edit_menu)


async def on_back_to_detail(_callback: CallbackQuery, _button: Button, manager: DialogManager):
    await manager.switch_to(SharedTestsSG.test_detail)


async def on_back_to_edit_menu(_callback: CallbackQuery, _button: Button, manager: DialogManager):
    await manager.switch_to(SharedTestsSG.edit_menu)


async def on_edit_password(_callback: CallbackQuery, _button: Button, manager: DialogManager):
    await manager.switch_to(SharedTestsSG.edit_password)


async def on_edit_attempts(_callback: CallbackQuery, _button: Button, manager: DialogManager):
    await manager.switch_to(SharedTestsSG.edit_attempts)


async def on_edit_time_limit(_callback: CallbackQuery, _button: Button, manager: DialogManager):
    await manager.switch_to(SharedTestsSG.edit_time_limit)


async def on_edit_group(_callback: CallbackQuery, _button: Button, manager: DialogManager):
    await manager.switch_to(SharedTestsSG.edit_group)


async def on_edit_expires(_callback: CallbackQuery, _button: Button, manager: DialogManager):
    await manager.switch_to(SharedTestsSG.edit_expires)


@inject
async def on_password_input(message: Message, _widget: MessageInput, manager: DialogManager, test_dao: FromDishka[TestDAO]):
    test_id = manager.dialog_data.get("selected_test_id")
    if not test_id:
        await message.answer("❌ Тест не найден")
        return
    
    if not message.text:
        await message.answer("❌ Пароль не может быть пустым")
        return
    
    password = message.text.strip()
    if len(password) > 255:
        await message.answer("❌ Пароль слишком длинный (максимум 255 символов)")
        return
    
    await test_dao.update(test_id, password=password)
    await message.answer("✅ Пароль обновлен")
    await manager.switch_to(SharedTestsSG.test_detail)


@inject
async def on_remove_password(_callback: CallbackQuery, _button: Button, manager: DialogManager, test_dao: FromDishka[TestDAO]):
    test_id = manager.dialog_data.get("selected_test_id")
    if not test_id:
        await _callback.answer("❌ Тест не найден")
        return
    
    await test_dao.update(test_id, password=None)
    await _callback.answer("✅ Пароль удален")
    await manager.switch_to(SharedTestsSG.test_detail)


@inject
async def on_attempts_input_edit(message: Message, _widget: MessageInput, manager: DialogManager, test_dao: FromDishka[TestDAO]):
    test_id = manager.dialog_data.get("selected_test_id")
    if not test_id:
        await message.answer("❌ Тест не найден")
        return
    
    if not message.text:
        await message.answer("❌ Количество попыток не может быть пустым")
        return
    
    attempts_str = message.text.strip()
    
    if not attempts_str.isdigit():
        await message.answer("❌ Количество попыток должно быть числом")
        return
    
    attempts = int(attempts_str)
    
    if attempts < 1:
        await message.answer("❌ Количество попыток должно быть больше 0")
        return
    
    if attempts > 100:
        await message.answer("❌ Количество попыток не может быть больше 100")
        return
    
    await test_dao.update(test_id, attempts=attempts)
    await message.answer("✅ Количество попыток обновлено")
    await manager.switch_to(SharedTestsSG.test_detail)


@inject
async def on_remove_attempts(_callback: CallbackQuery, _button: Button, manager: DialogManager, test_dao: FromDishka[TestDAO]):
    test_id = manager.dialog_data.get("selected_test_id")
    if not test_id:
        await _callback.answer("❌ Тест не найден")
        return
    
    await test_dao.update(test_id, attempts=None)
    await _callback.answer("✅ Ограничение попыток удалено")
    await manager.switch_to(SharedTestsSG.test_detail)


@inject
async def on_time_limit_input(message: Message, _widget: MessageInput, manager: DialogManager, test_dao: FromDishka[TestDAO]):
    test_id = manager.dialog_data.get("selected_test_id")
    if not test_id:
        await message.answer("❌ Тест не найден")
        return
    
    if not message.text:
        await message.answer("❌ Лимит времени не может быть пустым")
        return
    
    time_limit_str = message.text.strip()
    
    if not time_limit_str.isdigit():
        await message.answer("❌ Лимит времени должен быть числом (в минутах)")
        return
    
    time_limit_minutes = int(time_limit_str)
    
    if time_limit_minutes < 1:
        await message.answer("❌ Лимит времени должен быть больше 0")
        return
    
    if time_limit_minutes > 1440:
        await message.answer("❌ Лимит времени не может быть больше 1440 минут (24 часа)")
        return
    
    time_limit_seconds = time_limit_minutes * 60
    await test_dao.update(test_id, time_limit=time_limit_seconds)
    await message.answer("✅ Лимит времени обновлен")
    await manager.switch_to(SharedTestsSG.test_detail)


@inject
async def on_remove_time_limit(_callback: CallbackQuery, _button: Button, manager: DialogManager, test_dao: FromDishka[TestDAO]):
    test_id = manager.dialog_data.get("selected_test_id")
    if not test_id:
        await _callback.answer("❌ Тест не найден")
        return
    
    await test_dao.update(test_id, time_limit=None)
    await _callback.answer("✅ Лимит времени удален")
    await manager.switch_to(SharedTestsSG.test_detail)


@inject
async def get_groups_for_edit(dialog_manager: DialogManager, group_dao: FromDishka[GroupDAO], **_kwargs):
    groups = await group_dao.get_all()
    
    return {
        "groups": [(str(g.number), str(g.number)) for g in groups],
    }


@inject
async def on_group_selected_for_test(_callback: CallbackQuery, _widget, manager: DialogManager, item_id: str, test_dao: FromDishka[TestDAO]):
    test_id = manager.dialog_data.get("selected_test_id")
    if not test_id:
        await _callback.answer("❌ Тест не найден")
        return
    
    await test_dao.update(test_id, for_group=int(item_id))
    await _callback.answer("✅ Группа обновлена")
    await manager.switch_to(SharedTestsSG.test_detail)


@inject
async def on_remove_group(_callback: CallbackQuery, _button: Button, manager: DialogManager, test_dao: FromDishka[TestDAO]):
    test_id = manager.dialog_data.get("selected_test_id")
    if not test_id:
        await _callback.answer("❌ Тест не найден")
        return
    
    await test_dao.update(test_id, for_group=None)
    await _callback.answer("✅ Тест теперь доступен для всех групп")
    await manager.switch_to(SharedTestsSG.test_detail)


@inject
async def on_date_selected_for_test(_callback, _widget, manager: DialogManager, selected_date: date, test_dao: FromDishka[TestDAO]):
    test_id = manager.dialog_data.get("selected_test_id")
    if not test_id:
        await _callback.answer("❌ Тест не найден")
        return
    
    expires_at = datetime.combine(selected_date, time.min)
    await test_dao.update(test_id, expires_at=expires_at)
    await _callback.answer("✅ Срок действия обновлен")
    await manager.switch_to(SharedTestsSG.test_detail)


@inject
async def on_remove_expires(_callback: CallbackQuery, _button: Button, manager: DialogManager, test_dao: FromDishka[TestDAO]):
    test_id = manager.dialog_data.get("selected_test_id")
    if not test_id:
        await _callback.answer("❌ Тест не найден")
        return
    
    await test_dao.update(test_id, expires_at=None)
    await _callback.answer("✅ Срок действия удален")
    await manager.switch_to(SharedTestsSG.test_detail)


async def on_add_test_clicked(_callback: CallbackQuery, _button: Button, manager: DialogManager):
    await manager.start(SharedCreateTestSG.input_title, mode=StartMode.RESET_STACK)


async def on_back_clicked(_callback: CallbackQuery, _button: Button, manager: DialogManager):
    await manager.done()


@inject
async def get_groups_for_export(dialog_manager: DialogManager, group_dao: FromDishka[GroupDAO], **_kwargs):
    groups = await group_dao.get_all()
    return {
        "groups": [(f"🎓 {g.number}", str(g.number)) for g in groups],
        "count": len(groups),
    }


def create_excel_report(
    test_title: str,
    group_number: int,
    stats: list[tuple[str, int | None, datetime | None, bool | None]],
) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Статистика"
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    
    ws.merge_cells("A1:E1")
    ws["A1"] = f"Тест: {test_title}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    
    ws.merge_cells("A2:E2")
    ws["A2"] = f"Группа: {group_number}"
    ws["A2"].font = Font(bold=True, size=12)
    ws["A2"].alignment = Alignment(horizontal="center")
    
    headers = ["ФИО", "Результат (%)", "Оценка", "Дата прохождения", "Статус"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    passed_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    failed_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    not_passed_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    
    grades: list[int] = []
    
    for row_idx, (name, score, finished_at, is_passed) in enumerate(stats, 5):
        ws.cell(row=row_idx, column=1, value=name).border = thin_border
        
        if score is not None:
            ws.cell(row=row_idx, column=2, value=score).border = thin_border
            
            grade = score // 10
            grades.append(grade)
            ws.cell(row=row_idx, column=3, value=grade).border = thin_border
            
            finished_msk = to_msk(finished_at) if finished_at else None
            date_str = finished_msk.strftime("%d.%m.%Y %H:%M") if finished_msk else "—"
            ws.cell(row=row_idx, column=4, value=date_str).border = thin_border
            status = "Пройден" if is_passed else "Не пройден"
            status_cell = ws.cell(row=row_idx, column=5, value=status)
            status_cell.border = thin_border
            
            for col in range(1, 6):
                ws.cell(row=row_idx, column=col).fill = passed_fill if is_passed else failed_fill
        else:
            ws.cell(row=row_idx, column=2, value="—").border = thin_border
            ws.cell(row=row_idx, column=3, value="—").border = thin_border
            ws.cell(row=row_idx, column=4, value="—").border = thin_border
            status_cell = ws.cell(row=row_idx, column=5, value="Не проходил")
            status_cell.border = thin_border
            
            for col in range(1, 6):
                ws.cell(row=row_idx, column=col).fill = not_passed_fill
    
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 15
    
    total_users = len(stats)
    passed_users = sum(1 for _, score, _, is_passed in stats if score is not None and is_passed)
    attempted_users = sum(1 for _, score, _, _ in stats if score is not None)
    
    summary_row = len(stats) + 6
    ws.cell(row=summary_row, column=1, value="Итого:").font = Font(bold=True)
    ws.cell(row=summary_row + 1, column=1, value=f"Всего студентов: {total_users}")
    ws.cell(row=summary_row + 2, column=1, value=f"Прошли тест: {attempted_users}")
    ws.cell(row=summary_row + 3, column=1, value=f"Сдали: {passed_users}")
    if attempted_users > 0:
        success_rate = round(passed_users / attempted_users * 100)
        ws.cell(row=summary_row + 4, column=1, value=f"Процент сдачи: {success_rate}%")
    if grades:
        avg_grade = round(sum(grades) / len(grades), 1)
        ws.cell(row=summary_row + 5, column=1, value=f"Средняя оценка: {avg_grade}")
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


@inject
async def on_group_selected_for_export(
    _callback: CallbackQuery,
    _widget: Select,
    manager: DialogManager,
    item_id: str,
    test_dao: FromDishka[TestDAO],
    attempt_repo: FromDishka[TestAttemptRepository],
) -> None:
    test_id = manager.dialog_data.get("selected_test_id")
    if not test_id:
        await _callback.answer("❌ Тест не найден")
        return
    
    assert _callback.message is not None
    await _callback.answer("⏳ Формирую отчёт...")
    
    test = await test_dao.get_by_id(test_id)
    if not test:
        await _callback.message.answer("❌ Тест не найден")
        return
    
    group_number = int(item_id)
    stats = await attempt_repo.get_group_test_statistics(test_id, group_number)
    
    if not stats:
        await _callback.message.answer(f"❌ В группе {group_number} нет студентов")
        return
    
    excel_bytes = create_excel_report(test.title, group_number, stats)
    
    safe_title = "".join(c if c.isalnum() or c in "-_" else "_" for c in test.title)[:30]
    filename = f"{safe_title}_group_{group_number}.xlsx"
    
    await _callback.message.answer_document(
        document=BufferedInputFile(excel_bytes, filename=filename),
        caption=f"📊 <b>Статистика по тесту</b>\n\n📝 {test.title}\n🎓 Группа {group_number}",
    )


shared_tests_dialog = Dialog(
    Window(
        Format("<b>📝 Тесты</b>\n\nВсего: {count}"),
        ScrollingGroup(
            Select(
                Format("{item[0]}"),
                id="test_select",
                item_id_getter=lambda x: x[1],
                items="tests",
                on_click=on_test_selected,
            ),
            id="tests_scroll",
            width=1,
            height=7,
        ),
        Column(
            Button(Const("➕ Добавить тест"), id="add_test", on_click=on_add_test_clicked),
            Button(Const("◀️ Назад"), id="back", on_click=on_back_clicked),
        ),
        state=SharedTestsSG.tests_list,
        getter=get_tests_data,
    ),
    Window(
        Format("{test_info}"),
        Column(
            Button(
                Format("{button_text}"),
                id="toggle_active",
                on_click=on_toggle_active
            ),
            Button(
                Format("{results_button_text}"),
                id="toggle_results",
                on_click=on_toggle_results_viewable
            ),
            Button(Const("📊 Статистика"), id="statistics", on_click=on_statistics),
            Button(Const("🔗 Поделиться"), id="share", on_click=on_share_test),
            Button(Const("📥 Экспорт"), id="export", on_click=on_export_stats),
            Button(Const("✏️ Изменить"), id="edit_menu", on_click=on_edit_menu),
            Button(Const("◀️ Назад"), id="back", on_click=on_back_to_list),
        ),
        state=SharedTestsSG.test_detail,
        getter=get_test_detail,
    ),
    Window(
        Const("<b>✏️ Изменить тест</b>\n\nВыберите, что хотите изменить:"),
        Column(
            Button(Const("🔑 Пароль"), id="edit_password", on_click=on_edit_password),
            Button(Const("🔄 Попытки"), id="edit_attempts", on_click=on_edit_attempts),
            Button(Const("⏱️ Лимит времени"), id="edit_time_limit", on_click=on_edit_time_limit),
            Button(Const("👥 Группа"), id="edit_group", on_click=on_edit_group),
            Button(Const("📅 Срок действия"), id="edit_expires", on_click=on_edit_expires),
            Button(Const("◀️ Назад"), id="back", on_click=on_back_to_detail),
        ),
        state=SharedTestsSG.edit_menu,
    ),
    Window(
        Const("<b>🔑 Изменение пароля</b>\n\n💬 <b>Введите новый пароль</b> или удалите текущий:\n<i>(максимум 255 символов)</i>"),
        MessageInput(on_password_input),
        Column(
            Button(Const("🗑 Удалить пароль"), id="remove_password", on_click=on_remove_password),
            Button(Const("◀️ Назад"), id="back", on_click=on_back_to_edit_menu),
        ),
        state=SharedTestsSG.edit_password,
    ),
    Window(
        Const("<b>🔄 Изменение количества попыток</b>\n\n🔢 <b>Введите новое количество попыток</b> (1-100) или удалите ограничение:"),
        MessageInput(on_attempts_input_edit),
        Column(
            Button(Const("🗑 Без ограничений"), id="remove_attempts", on_click=on_remove_attempts),
            Button(Const("◀️ Назад"), id="back", on_click=on_back_to_edit_menu),
        ),
        state=SharedTestsSG.edit_attempts,
    ),
    Window(
        Const("<b>⏱️ Изменение лимита времени</b>\n\n🔢 <b>Введите лимит времени в минутах</b> (1-1440) или удалите ограничение:"),
        MessageInput(on_time_limit_input),
        Column(
            Button(Const("🗑 Без лимита"), id="remove_time_limit", on_click=on_remove_time_limit),
            Button(Const("◀️ Назад"), id="back", on_click=on_back_to_edit_menu),
        ),
        state=SharedTestsSG.edit_time_limit,
    ),
    Window(
        Const("<b>👥 Изменение группы</b>\n\n🎓 <b>Выберите группу</b> или удалите привязку:"),
        ScrollingGroup(
            Select(
                Format("{item[1]}"),
                id="groups",
                item_id_getter=lambda x: x[0],
                items="groups",
                on_click=on_group_selected_for_test,
            ),
            id="groups_scroll",
            width=2,
            height=7,
        ),
        Column(
            Button(Const("🗑 Для всех групп"), id="remove_group", on_click=on_remove_group),
            Button(Const("◀️ Назад"), id="back", on_click=on_back_to_edit_menu),
        ),
        state=SharedTestsSG.edit_group,
        getter=get_groups_for_edit,
    ),
    Window(
        Const("<b>📅 Изменение срока действия</b>\n\n🗓 <b>Выберите новую дату</b> или удалите срок:"),
        Calendar(id="calendar", on_click=on_date_selected_for_test),
        Column(
            Button(Const("🗑 Удалить срок"), id="remove_expires", on_click=on_remove_expires),
            Button(Const("◀️ Назад"), id="back", on_click=on_back_to_edit_menu),
        ),
        state=SharedTestsSG.edit_expires,
    ),
    Window(
        Format("<b>📊 Статистика теста</b>\n\nПрошли тест: {count}"),
        ScrollingGroup(
            Select(
                Format("{item[0]}"),
                id="attempt_select",
                item_id_getter=lambda x: x[1],
                items="attempts",
                on_click=on_attempt_selected,
            ),
            id="attempts_scroll",
            width=1,
            height=7,
        ),
        Column(
            Button(Const("🔄 Обновить"), id="refresh", on_click=on_statistics),
            Button(Const("◀️ Назад"), id="back", on_click=on_back_to_detail),
        ),
        state=SharedTestsSG.statistics,
        getter=get_statistics_data,
    ),
    Window(
        Format("{attempt_info}"),
        Button(Const("◀️ Назад"), id="back", on_click=on_back_to_statistics),
        state=SharedTestsSG.attempt_detail,
        getter=get_attempt_detail,
    ),
    Window(
        Format("<b>📥 Экспорт статистики</b>\n\nВыберите группу для экспорта:\n\nВсего групп: {count}"),
        ScrollingGroup(
            Select(
                Format("{item[0]}"),
                id="export_group_select",
                item_id_getter=lambda x: x[1],
                items="groups",
                on_click=on_group_selected_for_export,
            ),
            id="export_groups_scroll",
            width=2,
            height=7,
        ),
        Button(Const("◀️ Назад"), id="back", on_click=on_back_to_detail),
        state=SharedTestsSG.export_select_group,
        getter=get_groups_for_export,
    ),
)
