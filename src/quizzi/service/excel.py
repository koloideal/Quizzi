import io
from dataclasses import dataclass
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from quizzi.infrastructure.database.dao.test import TestDAO
from quizzi.infrastructure.database.repo.test_attempt import TestAttemptRepository
from quizzi.infrastructure.utils.timezone import to_msk


@dataclass
class ExcelReportResult:
    success: bool
    data: bytes | None = None
    filename: str | None = None
    caption: str = ""


class ExcelService:
    def __init__(self, test_dao: TestDAO, attempt_repo: TestAttemptRepository) -> None:
        self._test_dao = test_dao
        self._attempt_repo = attempt_repo
    
    def create_test_report(
        self,
        test_title: str,
        group_number: int,
        stats: list[tuple[str, int | None, datetime | None, bool | None]],
    ) -> bytes:
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Статистика"
        
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        
        ws.merge_cells("A1:E1")
        ws["A1"] = f"Тест: {test_title}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].alignment = Alignment(horizontal="center")
        
        ws.merge_cells("A2:E2")
        ws["A2"] = f"Группа: {group_number}"
        ws["A2"].font = Font(bold=True, size=12)
        ws["A2"].alignment = Alignment(horizontal="center")
        
        headers = ["ФИО", "Результат (%)", "Оценка", "Дата прохождения", "Статус"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        passed_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        failed_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        not_passed_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        
        grades: list[int] = []
        
        for row_idx, (name, score, finished_at, is_passed) in enumerate(stats, 5):
            ws.cell(row=row_idx, column=1, value=name).border = thin_border
            
            if score is not None:
                ws.cell(row=row_idx, column=2, value=score).border = thin_border
                
                grade = score // 10
                grades.append(grade)
                ws.cell(row=row_idx, column=3, value=grade).border = thin_border
                
                finished_msk = to_msk(finished_at) if finished_at else None
                date_str = finished_msk.strftime("%d.%m.%Y %H:%M") if finished_msk else "—"
                ws.cell(row=row_idx, column=4, value=date_str).border = thin_border
                status = "Пройден" if is_passed else "Не пройден"
                status_cell = ws.cell(row=row_idx, column=5, value=status)
                status_cell.border = thin_border
                
                for col in range(1, 6):
                    ws.cell(row=row_idx, column=col).fill = passed_fill if is_passed else failed_fill
            else:
                ws.cell(row=row_idx, column=2, value="—").border = thin_border
                ws.cell(row=row_idx, column=3, value="—").border = thin_border
                ws.cell(row=row_idx, column=4, value="—").border = thin_border
                status_cell = ws.cell(row=row_idx, column=5, value="Не проходил")
                status_cell.border = thin_border
                
                for col in range(1, 6):
                    ws.cell(row=row_idx, column=col).fill = not_passed_fill
        
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 15
        ws.column_dimensions["C"].width = 10
        ws.column_dimensions["D"].width = 20
        ws.column_dimensions["E"].width = 15
        
        total_users = len(stats)
        passed_users = sum(1 for _, score, _, is_passed in stats if score is not None and is_passed)
        attempted_users = sum(1 for _, score, _, _ in stats if score is not None)
        
        summary_row = len(stats) + 6
        ws.cell(row=summary_row, column=1, value="Итого:").font = Font(bold=True)
        ws.cell(row=summary_row + 1, column=1, value=f"Всего студентов: {total_users}")
        ws.cell(row=summary_row + 2, column=1, value=f"Прошли тест: {attempted_users}")
        ws.cell(row=summary_row + 3, column=1, value=f"Сдали: {passed_users}")
        if attempted_users > 0:
            success_rate = round(passed_users / attempted_users * 100)
            ws.cell(row=summary_row + 4, column=1, value=f"Процент сдачи: {success_rate}%")
        if grades:
            avg_grade = round(sum(grades) / len(grades), 1)
            ws.cell(row=summary_row + 5, column=1, value=f"Средняя оценка: {avg_grade}")
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.read()
    
    async def generate_group_report(self, test_id: int, group_number: int) -> ExcelReportResult:
        test = await self._test_dao.get_by_id(test_id)
        if not test:
            return ExcelReportResult(success=False, caption="❌ Тест не найден")
        
        stats = await self._attempt_repo.get_group_test_statistics(test_id, group_number)
        if not stats:
            return ExcelReportResult(success=False, caption=f"❌ В группе {group_number} нет студентов")
        
        excel_bytes = self.create_test_report(test.title, group_number, stats)
        safe_title = "".join(c if c.isalnum() or c in "-_" else "_" for c in test.title)[:30]
        filename = f"{safe_title}_group_{group_number}.xlsx"
        
        return ExcelReportResult(
            success=True,
            data=excel_bytes,
            filename=filename,
            caption=f"📊 <b>Статистика по тесту</b>\n\n📝 {test.title}\n🎓 Группа {group_number}",
        )
